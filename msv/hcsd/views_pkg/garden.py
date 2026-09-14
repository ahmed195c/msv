"""
Garden Follow-up — independent, standalone tracking of park/garden site
reviews (rodent activity around manholes, outside areas, and buildings).

URL prefix : /garden/
Templates  : hcsd/garden_*.html
"""

import io
import logging
import os
from itertools import groupby

from django.contrib.auth.decorators import login_required
from django.core.paginator import Paginator
from django.http import HttpResponse, HttpResponseForbidden
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.views.decorators.http import require_POST

from ..models import GARDEN_INFESTATION_TYPE_CHOICES, GardenAreaReview, GardenVisit, GardenVisitPhoto
from .common import _can_admin, _can_data_entry, _can_garden_monitor, _can_rodent_control_monitor

logger = logging.getLogger(__name__)

GARDEN_NOTE_CHOICES = [
    'إصابة خارجية / Outside Infestation',
    'منهول مصاب / Infested Manhole',
    'عدد المباني المصاب / Number of Infested Buildings',
]


def _infestation_type_from_post(request):
    valid_codes = {code for code, _ in GARDEN_INFESTATION_TYPE_CHOICES}
    selected = [code for code in request.POST.getlist('infestation_type') if code in valid_codes]
    return ','.join(selected)


def _can_manage(user):
    return (
        _can_admin(user) or _can_data_entry(user)
        or _can_garden_monitor(user) or _can_rodent_control_monitor(user)
    )


def _current_period_start():
    return timezone.localdate().replace(day=1)


def _garden_area_reviews_map(area_names, period_start):
    """Return {area_name: GardenAreaReview} for this month, creating any
    missing rows (defaulting to not-reviewed) so every area always has one."""
    reviews = {
        r.area_name: r
        for r in GardenAreaReview.objects.filter(area_name__in=area_names, period_start=period_start)
    }
    for name in area_names:
        if name not in reviews:
            reviews[name], _ = GardenAreaReview.objects.get_or_create(
                area_name=name, period_start=period_start,
            )
    return reviews


@login_required
def garden_list(request):
    query = (request.GET.get('q') or '').strip()
    show_reviewed = request.GET.get('reviewed') == '1'
    page_number = request.GET.get('page') or 1

    # Ordered to match the source Excel sheet's row order (its "No" column),
    # not the model's default newest-first ordering.
    visits_qs = GardenVisit.objects.order_by('id')
    if query:
        visits_qs = visits_qs.filter(area_name__icontains=query)

    period_start = _current_period_start()
    area_names = list(visits_qs.values_list('area_name', flat=True).distinct())
    area_reviews = _garden_area_reviews_map(area_names, period_start)

    reviewed_areas = {name for name, r in area_reviews.items() if r.is_reviewed}
    pending_count  = len(area_reviews) - len(reviewed_areas)

    if show_reviewed:
        visits_qs = visits_qs.filter(area_name__in=reviewed_areas)
    else:
        visits_qs = visits_qs.exclude(area_name__in=reviewed_areas)

    paginator = Paginator(visits_qs, 20)
    page_obj = paginator.get_page(page_number)
    for row in page_obj:
        row.review = area_reviews.get(row.area_name)

    return render(request, 'hcsd/garden_list.html', {
        'rows': page_obj,
        'page_obj': page_obj,
        'query': query,
        'show_reviewed': show_reviewed,
        'pending_count': pending_count,
        'reviewed_count': len(reviewed_areas),
        'total_count': paginator.count,
        'can_manage': _can_manage(request.user),
        'can_admin': _can_admin(request.user),
    })


@login_required
def garden_create(request):
    if not _can_manage(request.user):
        return redirect('garden_list')

    errors = []
    if request.method == 'POST':
        area_name        = (request.POST.get('area_name') or '').strip()
        location_details = (request.POST.get('location_details') or '').strip()
        google_maps_url  = (request.POST.get('google_maps_url') or '').strip()

        def _int(name):
            raw = (request.POST.get(name) or '').strip()
            try:
                return int(raw) if raw else None
            except ValueError:
                return None

        def _float(name):
            raw = (request.POST.get(name) or '').strip()
            try:
                return float(raw) if raw else None
            except ValueError:
                return None

        if not area_name:
            errors.append('يرجى إدخال اسم المنطقة.')

        if not errors:
            obj = GardenVisit.objects.create(
                visit_date=timezone.localdate(),
                area_name=area_name,
                location_details=location_details,
                google_maps_url=google_maps_url,
                latitude=_float('latitude'),
                longitude=_float('longitude'),
                infested_manholes=_int('infested_manholes'),
                infested_outside=_int('infested_outside'),
                total_infested_bldg=_int('total_infested_bldg'),
                infestation_type=_infestation_type_from_post(request),
                created_by=request.user,
            )

            # Each "spot_photos_<N>" / "spot_description_<N>" pair is one
            # infestation spot — a visit can have several, each with its
            # own set of photos and an optional shared description.
            spot_indices = set()
            for key in request.POST:
                if key.startswith('spot_description_'):
                    spot_indices.add(key[len('spot_description_'):])
            for key in request.FILES:
                if key.startswith('spot_photos_'):
                    spot_indices.add(key[len('spot_photos_'):])

            for idx in spot_indices:
                try:
                    spot_number = int(idx)
                except ValueError:
                    continue
                description = (request.POST.get(f'spot_description_{idx}') or '').strip()
                for photo in request.FILES.getlist(f'spot_photos_{idx}'):
                    GardenVisitPhoto.objects.create(
                        garden_visit=obj, spot_number=spot_number,
                        description=description, file=photo, uploaded_by=request.user,
                    )
            return redirect('garden_list')

    return render(request, 'hcsd/garden_create.html', {
        'errors': errors,
        'post': request.POST,
        'infestation_type_choices': GARDEN_INFESTATION_TYPE_CHOICES,
        'selected_infestation_types': request.POST.getlist('infestation_type'),
    })


@login_required
def garden_detail(request, pk):
    obj = get_object_or_404(GardenVisit, pk=pk)
    can_manage = _can_manage(request.user)

    if request.method == 'POST':
        if not can_manage:
            return HttpResponseForbidden()

        def _int(name):
            raw = (request.POST.get(name) or '').strip()
            try:
                return int(raw) if raw else None
            except ValueError:
                return None

        def _float(name):
            raw = (request.POST.get(name) or '').strip()
            try:
                return float(raw) if raw else None
            except ValueError:
                return None

        if request.POST.get('action') == 'save_location':
            obj.latitude         = _float('latitude')
            obj.longitude        = _float('longitude')
            obj.google_maps_url  = (request.POST.get('google_maps_url') or '').strip()
            obj.location_details = (request.POST.get('location_details') or '').strip()
            obj.save(update_fields=['latitude', 'longitude', 'google_maps_url', 'location_details'])
        else:
            obj.infested_manholes    = _int('infested_manholes')
            obj.infested_outside     = _int('infested_outside')
            obj.total_infested_bldg  = _int('total_infested_bldg')
            obj.notes                = (request.POST.get('notes') or '').strip()
            obj.infestation_type     = _infestation_type_from_post(request)
            obj.save(update_fields=[
                'infested_manholes', 'infested_outside', 'total_infested_bldg',
                'notes', 'infestation_type',
            ])

            # Updating the infestation data for an area counts as reviewing
            # it for this month — no need for a separate manual step.
            review, _created = GardenAreaReview.objects.get_or_create(
                area_name=obj.area_name, period_start=_current_period_start(),
            )
            if not review.is_reviewed:
                review.is_reviewed = True
                review.reviewed_by = request.user
                review.reviewed_at = timezone.now()
                review.save(update_fields=['is_reviewed', 'reviewed_by', 'reviewed_at'])
        return redirect('garden_detail', pk=pk)

    photo_spots = []
    for spot_number, group in groupby(obj.photos.all(), key=lambda p: p.spot_number):
        group = list(group)
        photo_spots.append({
            'spot_number': spot_number,
            'description': group[0].description,
            'photos': group,
        })

    return render(request, 'hcsd/garden_detail.html', {
        'obj': obj,
        'can_manage': can_manage,
        'can_admin': _can_admin(request.user),
        'note_choices': GARDEN_NOTE_CHOICES,
        'infestation_type_choices': GARDEN_INFESTATION_TYPE_CHOICES,
        'photo_spots': photo_spots,
    })


@login_required
def garden_visit_report(request, pk):
    """Download a single Word document with a visit's full details and all
    of its infestation-spot photos — for archiving one request at a time."""
    obj = get_object_or_404(GardenVisit, pk=pk)

    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.shared import Inches, Pt, RGBColor
    except ImportError:
        return HttpResponse('مكتبة python-docx غير مثبتة.', status=500)

    doc = Document()

    section = doc.sections[0]
    section.page_width  = int(8.27 * 914400)
    section.page_height = int(11.69 * 914400)
    section.left_margin = section.right_margin = int(1 * 914400)
    section.top_margin  = section.bottom_margin = int(1 * 914400)

    def set_rtl(paragraph):
        pPr = paragraph._p.get_or_add_pPr()
        bidi = OxmlElement('w:bidi')
        pPr.append(bidi)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    def add_heading(text):
        p = doc.add_paragraph()
        set_rtl(p)
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(14)
        run.font.color.rgb = RGBColor(0x1a, 0x3a, 0x5c)
        return p

    def add_field(label, value):
        p = doc.add_paragraph()
        set_rtl(p)
        lbl = p.add_run(f'{label}: ')
        lbl.bold = True
        lbl.font.size = Pt(11)
        val = p.add_run(str(value) if value not in (None, '') else '—')
        val.font.size = Pt(11)

    def add_section_title(text):
        p = doc.add_paragraph()
        set_rtl(p)
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(0x2a, 0x6a, 0x9a)

    title = doc.add_paragraph()
    set_rtl(title)
    t = title.add_run('تقرير زيارة متابعة المناطق')
    t.bold = True
    t.font.size = Pt(20)
    t.font.color.rgb = RGBColor(0x1a, 0x3a, 0x5c)
    doc.add_paragraph()

    add_heading('بيانات الموقع')
    add_field('اسم المنطقة', obj.area_name)
    add_field('تفاصيل الموقع', obj.location_details)
    add_field('تاريخ الزيارة', obj.visit_date.strftime('%d/%m/%Y') if obj.visit_date else None)
    if obj.latitude is not None and obj.longitude is not None:
        add_field('الإحداثيات', f'{obj.latitude}, {obj.longitude}')
        add_field(
            'رابط خرائط قوقل',
            obj.google_maps_url or f'https://www.google.com/maps?q={obj.latitude},{obj.longitude}',
        )
    elif obj.google_maps_url:
        add_field('رابط خرائط قوقل', obj.google_maps_url)
    doc.add_paragraph()

    add_heading('بيانات الإصابة')
    add_field('مناهيل مصابة', obj.infested_manholes)
    add_field('إصابة خارجية', obj.infested_outside)
    add_field('إجمالي مباني مصابة', obj.total_infested_bldg)
    if obj.infestation_type:
        add_field('نوع الإصابة', obj.infestation_type_display)
    if obj.notes:
        add_field('ملاحظات', obj.notes)
    doc.add_paragraph()

    # Photos are re-encoded to a capped resolution/quality before embedding —
    # original phone-camera photos can be several MB each.
    from PIL import Image, ImageOps
    _PHOTO_MAX_DIM = 1000
    _PHOTO_JPEG_QUALITY = 70

    photos_qs = obj.photos.all()
    if photos_qs:
        add_heading('أماكن الإصابة')
        for spot_number, group in groupby(photos_qs, key=lambda p: p.spot_number):
            group = list(group)
            add_section_title(f'مكان الإصابة {spot_number}')
            if group[0].description:
                desc = doc.add_paragraph(group[0].description)
                set_rtl(desc)
            for photo in group:
                try:
                    photo_path = photo.file.path
                    if not os.path.exists(photo_path):
                        continue
                    with Image.open(photo_path) as img:
                        img = ImageOps.exif_transpose(img).convert('RGB')
                        img.thumbnail((_PHOTO_MAX_DIM, _PHOTO_MAX_DIM), Image.LANCZOS)
                        photo_buf = io.BytesIO()
                        img.save(photo_buf, format='JPEG', quality=_PHOTO_JPEG_QUALITY, optimize=True)
                        photo_buf.seek(0)
                    doc.add_picture(photo_buf, width=Inches(4.5))
                except Exception:
                    logger.exception('Failed to add photo %s to garden visit report', photo.pk)
            doc.add_paragraph()

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    filename = f'garden_visit_{obj.pk}.docx'
    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


@login_required
@require_POST
def garden_delete(request, pk):
    obj = get_object_or_404(GardenVisit, pk=pk)
    if not _can_admin(request.user):
        return HttpResponseForbidden()

    obj.delete()
    return redirect('garden_list')


@login_required
@require_POST
def garden_area_review_toggle(request, pk):
    review = get_object_or_404(GardenAreaReview, pk=pk)
    if not _can_manage(request.user):
        return HttpResponseForbidden()

    review.is_reviewed = not review.is_reviewed
    review.reviewed_by = request.user if review.is_reviewed else None
    review.reviewed_at = timezone.now() if review.is_reviewed else None
    review.save(update_fields=['is_reviewed', 'reviewed_by', 'reviewed_at'])

    next_url = request.POST.get('next') or ''
    if next_url.startswith('/garden/'):
        return redirect(next_url)
    return redirect('garden_list')

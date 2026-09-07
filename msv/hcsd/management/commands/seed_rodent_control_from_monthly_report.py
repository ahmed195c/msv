"""
Historical import from "MONTHLY REPORT RODENT CONTROL 2026.xlsx" — the
actual field-visit log the real "Rodent Team Report" export is modeled on.
Unlike Park RBS and Manholes Report.xlsx (per-site, one row per building
per month), this file is per-visit: an area can have several rows in the
same month, each a separate team's visit, and it already carries every
field the new report needs (team leader, technicians, time in/out, and
all the inspected/infested counts by category).

Safe to re-run — visits are matched on
(building, period_start, visit_date, time_in, team_leader_name), so
re-running after fixing the source file updates existing rows instead of
duplicating them.
"""
import datetime
import re

import openpyxl
from django.core.management.base import BaseCommand, CommandError

from hcsd.models import RodentControlBuilding, RodentControlVisit

DEFAULT_PATH = 'hcsd/static/hcsd/excl/buldings/MONTHLY REPORT RODENT CONTROL 2026.xlsx'

MONTH_SHEETS = {
    'Jan': (2026, 1), 'FEB': (2026, 2), 'MAR': (2026, 3), 'APR': (2026, 4),
    'MAY': (2026, 5), 'JUNE': (2026, 6), 'JULY': (2026, 7), 'AUG': (2026, 8),
    'SEP': (2026, 9),
}

# Header text -> model field. A couple of sheets have typos in the source
# ("Invested" for "Infested") — both spellings map to the same field.
FIELD_MAP = {
    'Team Leader': 'team_leader_name',
    'No of Technicians': 'technicians_count',
    'Total number of Inspecteted Bldg/ Villa': 'bldg_villa_inspected_count',
    'Number of Infested Bldg/ Villa': 'bldg_villa_infested_count',
    'Number of Invested Bldg/ Villa': 'bldg_villa_infested_count',
    'Total Number of  Manholes': 'manholes_inspected_count',
    'Number of Insp & Treated Manholes': 'manholes_treated_count',
    ' Number of Infested Manholes': 'manholes_infested_count',
    'Total number of Inspected Burrows': 'burrows_outside_count',
    'Number of Infested Burrows': 'burrows_infested_count',
    'Total Inspected Construction': 'construction_inspected_count',
    ' Infested Construction': 'construction_infested_count',
    'Totol InspectedMasjed': 'masjed_inspected_count',
    ' InfestedMasjed': 'masjed_infested_count',
    'Total Inspected Trees': 'trees_inspected_count',
    ' Insfested Trees': 'trees_infested_count',
    'Total Inspected Electrical Rooms': 'electrical_inspected_count',
    ' Insfested Electrical Rooms': 'electrical_infested_count',
    'Gov Offices': 'gov_offices_count',
}

PRODUCT_KEYWORDS = [
    ('SUREFIRE', 'rodenticide_surefire_qty'),
    ('FACORAT', 'rodenticide_facorat_qty'),
    ('VERTOX', 'rodenticide_vertox_qty'),
    ('SELLIOX', 'rodenticide_sellioxid_qty'),
    ('VICTOR', 'rodenticide_victor_qty'),
    ('PROTECT', 'rodenticide_protect_qty'),
    ('SENSATION', 'rodenticide_protect_qty'),
    ('NOCURAT', 'rodenticide_nocurat_qty'),
]


def _rodenticide_field(header):
    if '{' in header and '}' in header:
        name = header.split('{', 1)[1].split('}', 1)[0]
    else:
        m = re.match(r'Rodenticide\s+\d+\s+(.*)', header)
        if not m:
            return None
        name = re.sub(r'\s*Qyt$', '', m.group(1), flags=re.IGNORECASE)
    name = name.upper()
    for keyword, field in PRODUCT_KEYWORDS:
        if keyword in name:
            return field
    return None


def _find_header_row(ws):
    for r in range(1, 5):
        vals = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
        if 'Team Leader' in vals and 'Area Name' in vals:
            return r
    return None


class Command(BaseCommand):
    help = 'One-time import of historical rodent-control data from MONTHLY REPORT RODENT CONTROL 2026.xlsx'

    def add_arguments(self, parser):
        parser.add_argument('--path', default=DEFAULT_PATH)
        parser.add_argument('--dry-run', action='store_true')

    def handle(self, *args, **options):
        path = options['path']
        dry_run = options['dry_run']

        try:
            wb = openpyxl.load_workbook(path, data_only=True)
        except FileNotFoundError:
            raise CommandError(f'File not found: {path}')

        buildings_created = 0
        visits_created = 0
        visits_updated = 0
        rows_skipped = 0
        skipped_sheets = []

        for sheet_name, (year, month) in MONTH_SHEETS.items():
            if sheet_name not in wb.sheetnames:
                skipped_sheets.append(sheet_name)
                continue
            ws = wb[sheet_name]
            header_row = _find_header_row(ws)
            if header_row is None:
                skipped_sheets.append(sheet_name)
                continue

            headers = [ws.cell(row=header_row, column=c).value for c in range(1, ws.max_column + 1)]
            col_of = {h: i + 1 for i, h in enumerate(headers) if h}

            area_col = col_of.get('Area Name')
            leader_col = col_of.get('Team Leader')
            date_col = col_of.get('Date')
            time_in_col = col_of.get('Time In')
            time_out_col = col_of.get('Time Out:')
            if not area_col:
                skipped_sheets.append(sheet_name)
                continue

            field_cols = {col_of[h]: f for h, f in FIELD_MAP.items() if h in col_of}
            rodenticide_cols = {}
            for h, c in col_of.items():
                if h and 'Rodenticide' in h:
                    field = _rodenticide_field(h)
                    if field:
                        rodenticide_cols[c] = field

            period_start = datetime.date(year, month, 1)

            def _cell_int(row, col):
                if not col:
                    return None
                raw = ws.cell(row=row, column=col).value
                if raw is None or (isinstance(raw, str) and not raw.strip()):
                    return None
                try:
                    return int(float(raw))
                except (TypeError, ValueError):
                    return None

            def _cell_time(row, col):
                if not col:
                    return None
                raw = ws.cell(row=row, column=col).value
                if isinstance(raw, datetime.time):
                    return raw
                if isinstance(raw, datetime.datetime):
                    return raw.time()
                return None

            for r in range(header_row + 1, ws.max_row + 1):
                name = ws.cell(row=r, column=area_col).value
                if not name or not str(name).strip():
                    continue
                name = str(name).strip()
                if name.upper() == 'TOTAL':
                    continue

                raw_date = ws.cell(row=r, column=date_col).value if date_col else None
                visit_date = raw_date.date() if isinstance(raw_date, datetime.datetime) else (
                    raw_date if isinstance(raw_date, datetime.date) else None
                )

                team_leader_name = ''
                if leader_col:
                    raw_leader = ws.cell(row=r, column=leader_col).value
                    team_leader_name = str(raw_leader).strip() if raw_leader else ''

                # A row with no visit_date and no team leader is a blank
                # trailing row past the real data — skip it.
                if not visit_date and not team_leader_name:
                    rows_skipped += 1
                    continue

                if dry_run:
                    building = RodentControlBuilding.objects.filter(name__iexact=name).first()
                else:
                    building = RodentControlBuilding.objects.filter(name__iexact=name).first()
                    if not building:
                        building = RodentControlBuilding.objects.create(name=name, area=name)
                        buildings_created += 1

                defaults = {'team_leader_name': team_leader_name}
                for col, field in field_cols.items():
                    if field == 'team_leader_name':
                        continue
                    defaults[field] = _cell_int(r, col)
                for col, field in rodenticide_cols.items():
                    v = _cell_int(r, col)
                    if v is not None:
                        defaults[field] = defaults.get(field, 0) + v

                if dry_run:
                    continue
                if not building:
                    continue

                visit, created = RodentControlVisit.objects.update_or_create(
                    building=building, period_start=period_start,
                    visit_date=visit_date, time_in=_cell_time(r, time_in_col),
                    team_leader_name=team_leader_name,
                    defaults={**defaults, 'time_out': _cell_time(r, time_out_col)},
                )
                if created:
                    visits_created += 1
                else:
                    visits_updated += 1

        self.stdout.write(self.style.SUCCESS(
            f'Buildings created: {buildings_created}. '
            f'Visits created: {visits_created}, updated: {visits_updated}. '
            f'Blank rows skipped: {rows_skipped}. '
            f'Skipped sheets: {skipped_sheets or "none"}.'
        ))
